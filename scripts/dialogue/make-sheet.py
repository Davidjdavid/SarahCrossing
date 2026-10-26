"""The dialogue spreadsheet from dialogue-lines.json (export.mjs): a Read me, every line in the game, the new lines for
the Town Hall, Post Office, Bank and zip code, and who says how much.

    python3 scripts/dialogue/make-sheet.py docs/dialogue/dialogue-lines.json docs/dialogue/dialogue.xlsx
"""
import json
import sys
from collections import Counter

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

src, out = sys.argv[1], sys.argv[2]
data = json.load(open(src))
rows, fresh = data['rows'], data['fresh']

HEAD = Font(bold=True, color='FFFFFF')
HEAD_FILL = PatternFill('solid', fgColor='4A6B3A')
EDIT_FILL = PatternFill('solid', fgColor='FFF6D6')
WRAP = Alignment(wrap_text=True, vertical='top')
TOP = Alignment(vertical='top')
KIND_FILL = {
    'Dialogue': 'FFFFFF', 'Narration': 'EEF3FA', 'Screen text': 'F2F2F2', 'Button or label': 'F2F2F2', 'Choice': 'F6EEFA',
    'Message': 'F2F2F2', 'Letter': 'FAF1E8', 'Text message': 'EAF6EA',
}
STATE_FILL = {
    'Works': 'FFFFFF', 'Loading': 'EEF3FA', 'No internet': 'FBE9E7', 'Service down': 'FBE9E7', 'Old data': 'FFF3E0',
    'No zip code': 'FFF3E0', 'Bad zip code': 'FFF3E0', 'Market closed': 'EDE7F6', 'Not enough money': 'FFF3E0', 'Empty': 'F2F2F2',
}


def table(ws, heads, widths, data_rows, wrap_cols, edit_cols, fill_of=None):
    ws.append(heads)
    for c, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(c)].width = w
        cell = ws.cell(row=1, column=c)
        cell.font, cell.fill, cell.alignment = HEAD, HEAD_FILL, Alignment(vertical='center', wrap_text=True)
    ws.row_dimensions[1].height = 30
    for r in data_rows:
        ws.append(r)
        i = ws.max_row
        fill = PatternFill('solid', fgColor=fill_of(r)) if fill_of else None
        for c in range(1, len(heads) + 1):
            cell = ws.cell(row=i, column=c)
            cell.alignment = WRAP if c in wrap_cols else TOP
            if c in edit_cols:
                cell.fill = EDIT_FILL
            elif fill:
                cell.fill = fill
    ws.freeze_panes = 'B2'
    ws.auto_filter.ref = f'A1:{get_column_letter(len(heads))}{ws.max_row}'


wb = Workbook()

# ---------------------------------------------------------------- read me
ws = wb.active
ws.title = 'Read me'
ws.column_dimensions['A'].width = 110
notes = [
    ('Sarah Crossing: every line of dialogue', Font(bold=True, size=16)),
    (f'Made {data["made"]}. {len(rows):,} lines in the game today, plus {len(fresh)} new lines for the features to come.', None),
    ('', None),
    ('How to edit', Font(bold=True, size=13)),
    ('1. Change a line by writing the new wording in the yellow "Your new line" column. Leave it empty to keep the line as it is.', None),
    ('2. Notes for me (who should say it, when it should come up, "cut this", "write more like this") go in the yellow "Notes" column.', None),
    ('3. Please don\'t change the ID, Line or Original columns: they\'re how your changes find their way back into the game.', None),
    ('4. Words in {curly brackets} are filled in by the game: {name} is Sarah\'s name, {who} a neighbor\'s, {island} the island\'s,', None),
    ('   and {something(like this)} is worked out by the code (a number, an item\'s name). Keep them in your new line where they belong.', None),
    ('5. " | " in a line separates choices or versions shown one at a time.', None),
    ('6. Send the file back when you\'re done (even partly done) and I\'ll put the changes in.', None),
    ('', None),
    ('The sheets', Font(bold=True, size=13)),
    ('Dialogue: everything people say, the narration, letters, texts and Sarah\'s choices, in the order the code has them', None),
    ('   (neighbors\' lines are grouped by kind: everyday chat, the hour, the weather, where they are, what Sarah\'s wearing,', None),
    ('   holidays...). Use the filter arrows on the top row to see one person, one kind of line or one part of the game at a time.', None),
    ('Screen text: menus, buttons, labels, item and place names, and the little messages that pop up. Edit these the same way.', None),
    ('New - Town Hall, PO, Bank: lines for the zip code, the forecast, the newspaper and the three buildings, written', None),
    ('   for every way each thing can go: it worked, still loading, no internet, the news / weather / stock service is down,', None),
    ('   no zip code yet, the market is closed, not enough money, nothing there yet. Nothing here is in the game yet.', None),
    ('   Who works there is my proposal for you to change: the Cozy Council (Sir Floret, Dame Stalk, Little Sprout) runs the', None),
    ('   Town Hall; the Post Office and the Bank have new keepers, Pita (a pita pocket full of letters) and Penny (a chocolate', None),
    ('   coin in gold foil). Rename them, or say who should do it instead.', None),
    ('People: how many lines each person has.', None),
    ('', None),
    ('Columns', Font(bold=True, size=13)),
    ('Who: who says it. "Narration" is the game describing something; "On screen" is a menu, label or message; "Any neighbor"', None),
    ('   is said by whichever neighbor it comes up with; "The neighbor Sarah is talking to" is the same, picked by the code.', None),
    ('   "A" and "B" are the two neighbors in a drama storyline. "Sarah" lines are her choices in a conversation.', None),
    ('To: who it\'s said to, when it isn\'t Sarah (a visitor playing as David, James, John, Joe or Elizabeth).', None),
    ('When: when it comes up, in plain words.', None),
    ('Where: the part of the game.', None),
    ('Kind: dialogue, narration, screen text, a choice, a message that pops up, a letter, a text message.', None),
    ('Source: the file and line in the code (for me).', None),
    ('', None),
    ('Nothing anyone says has been changed: this is the game as it is today.', Font(italic=True)),
]
for text, font in notes:
    ws.append([text])
    if font:
        ws.cell(row=ws.max_row, column=1).font = font

# ---------------------------------------------------------------- every line: what people say, then what's on screen
SCREEN = {'Screen text', 'Button or label', 'Message'}
heads = ['ID', 'Who', 'To', 'Line', 'When', 'Where', 'Kind', 'Your new line', 'Notes', 'Source', 'Original (don\'t change)']
widths = [17, 18, 11, 70, 46, 22, 14, 50, 30, 20, 12]
for title, pick in (('Dialogue', lambda r: r['kind'] not in SCREEN), ('Screen text', lambda r: r['kind'] in SCREEN)):
    ws = wb.create_sheet(title)
    table(ws, heads, widths,
          [[r['id'], r['who'], r['to'], r['line'], r['when'], r['area'], r['kind'], '', '', r['src'], r['raw']] for r in rows if pick(r)],
          wrap_cols={4, 5, 8, 9}, edit_cols={8, 9}, fill_of=lambda r: KIND_FILL.get(r[6], 'FFFFFF'))
    ws.column_dimensions['K'].hidden = True

# ---------------------------------------------------------------- the new lines
ws = wb.create_sheet('New - Town Hall, PO, Bank')
heads = ['ID', 'Feature', 'Who', 'How it went', 'When', 'Line', 'Your new line', 'Notes']
widths = [15, 18, 16, 15, 46, 70, 50, 30]
table(ws, heads, widths,
      [[l['id'], l['feature'], l['who'], l['state'], l['when'], l['line'], '', ''] for l in fresh],
      wrap_cols={5, 6, 7, 8}, edit_cols={7, 8}, fill_of=lambda r: STATE_FILL.get(r[3], 'FFFFFF'))

# ---------------------------------------------------------------- who says how much
ws = wb.create_sheet('People')
count = Counter(r['who'] for r in rows)
new_count = Counter(l['who'] for l in fresh)
table(ws, ['Who', 'Lines in the game', 'New lines'], [32, 18, 12],
      [[w, count.get(w, 0), new_count.get(w, 0)] for w in sorted(set(count) | set(new_count), key=lambda w: -(count.get(w, 0) + new_count.get(w, 0)))],
      wrap_cols=set(), edit_cols=set())

wb.save(out)
