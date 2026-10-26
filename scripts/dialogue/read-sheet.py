"""The edits in a dialogue spreadsheet David sent back, as JSON for apply.mjs: each changed line (and each note), and
his edits to the new lines.

    python3 scripts/dialogue/read-sheet.py dialogue.xlsx edits.json
"""
import json
import sys

from openpyxl import load_workbook

wb = load_workbook(sys.argv[1], read_only=True)
out = {'lines': [], 'fresh': []}
ORIGINAL = "Original (don't change)"


def rows(sheet):
    head = None
    for row in wb[sheet].iter_rows(values_only=True):
        if head is None:
            head = list(row)
            continue
        yield dict(zip(head, row))


for sheet in ('Dialogue', 'Screen text'):
    if sheet not in wb.sheetnames:
        continue
    for r in rows(sheet):
        new, notes = r.get('Your new line'), r.get('Notes') or ''
        if new not in (None, '') or notes:
            out['lines'].append({'id': r['ID'], 'new': None if new in (None, '') else str(new), 'src': r['Source'], 'raw': r[ORIGINAL], 'line': r['Line'], 'notes': notes})
for sheet in wb.sheetnames:
    if sheet.startswith('New'):
        for r in rows(sheet):
            if r.get('Your new line') or r.get('Notes'):
                out['fresh'].append({'id': r['ID'], 'new': r.get('Your new line'), 'line': r['Line'], 'notes': r.get('Notes') or ''})
json.dump(out, open(sys.argv[2], 'w'))
